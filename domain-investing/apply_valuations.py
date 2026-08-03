"""Enter sourced valuations. Every value here must name where it came from."""
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment

WB = "/home/user/AlphaEdge/domain-investing/domain-portfolio-tracker.xlsx"
FONT, BLUE = "Arial", "0000FF"

# domain -> (value, source). Brokerage figure used: it is the optimistic end of HumbleWorth's
# three estimates, so if a name fails at this level it fails at every level.
VALUATIONS = {
    "brewingfinancing.com": (285, "HumbleWorth brokerage 2026-08-03 (mktpl $59, auction $0)"),
}

wb = load_workbook(WB)
ws = wb["Candidates"]
index = {ws[f"A{r}"].value: r for r in range(5, 205) if ws[f"A{r}"].value}
missing = [d for d in VALUATIONS if d not in index]
assert not missing, f"not on the sheet: {missing}"

for dom, (val, src) in VALUATIONS.items():
    r = index[dom]
    ws[f"K{r}"].value = val
    ws[f"K{r}"].number_format = '$#,##0;($#,##0);-'
    ws[f"L{r}"].value = src
    for col in ("K", "L"):
        ws[f"{col}{r}"].font = Font(name=FONT, size=10, color=BLUE)
    ws[f"K{r}"].alignment = Alignment(horizontal="center", vertical="center")
    ws[f"L{r}"].alignment = Alignment(horizontal="left", vertical="center")

wb.save(WB)
print(f"entered {len(VALUATIONS)} sourced valuation(s)")
