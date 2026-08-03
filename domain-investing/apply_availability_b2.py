"""Apply GoDaddy bulk-search results for batch 2 (2026-08-03), rows 61-112."""

from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment

WB = "/home/user/AlphaEdge/domain-investing/domain-portfolio-tracker.xlsx"
FONT, BLUE = "Arial", "0000FF"
MONEY2 = '$#,##0.00;($#,##0.00);-'
FIRST, LAST = 61, 112
HANDREG = 22.99

# 8 of 52 available. Only stipsheet.com is at registration price.
AVAILABLE = {
    "stipsheet.com": 22.99,
    "ticketsize.com": 2395.00,
    "loantape.com": 2988.00,
    "approvefast.com": 3495.00,
    "residualvalue.com": 14499.00,
    "fundpilot.com": 14999.00,
    "originations.com": 55000.00,
    "servicing.com": 287500.00,
}

wb = load_workbook(WB)
ws = wb["Candidates"]

seen, n_avail = [], 0
for r in range(FIRST, LAST + 1):
    dom = ws[f"A{r}"].value
    if not dom:
        continue
    seen.append(dom)
    if dom in AVAILABLE:
        n_avail += 1
        price = AVAILABLE[dom]
        ws[f"J{r}"].value = price
        ws[f"J{r}"].number_format = MONEY2
        status = "Available - reg" if price == HANDREG else "Available - premium"
    else:
        ws[f"J{r}"].value = None
        status = "Taken"
    ws[f"W{r}"].value = status
    for col in ("J", "W"):
        c = ws[f"{col}{r}"]
        c.font = Font(name=FONT, size=10, color=BLUE)
        c.alignment = Alignment(horizontal="center", vertical="center")

missing = [d for d in AVAILABLE if d not in seen]
assert not missing, f"available names not found in batch 2 rows: {missing}"

ws["A2"] = ("Niche 1: equipment financing & commercial lending, both batches checked at GoDaddy "
            "2026-08-03. Batch 1 rows 5-60: 20 of 56 available. Batch 2 rows 61-112: 8 of 52 "
            "available, only one at registration price. Across all 108 names, every name available "
            "at registration price scores below the BUY threshold. Premium figures are ASKING "
            "prices, not sale comps.")
ws["A2"].font = Font(name=FONT, size=10, italic=True)

wb.save(WB)
print(f"batch 2: checked {len(seen)} rows | available {n_avail} | taken {len(seen)-n_avail}")
