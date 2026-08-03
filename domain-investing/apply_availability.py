"""Apply GoDaddy bulk-search results (2026-08-03) to the Candidates tab."""

from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment

WB = "/home/user/AlphaEdge/domain-investing/domain-portfolio-tracker.xlsx"
FONT, BLUE = "Arial", "0000FF"
MONEY2 = '$#,##0.00;($#,##0.00);-'

# 20 available, with GoDaddy's asking price. $22.99 = standard registration.
AVAILABLE = {
    "torquelending.com": 22.99,
    "agequipmentloans.com": 22.99,
    "section179financing.com": 22.99,
    "minefinancing.com": 22.99,
    "printfinancing.com": 22.99,
    "brewingfinancing.com": 22.99,
    "kitchenfinancing.com": 2788.00,
    "sba504.com": 2795.00,
    "fundrig.com": 2995.00,
    "commercial-lending.com": 3595.00,
    "trailerfinancing.com": 4888.00,
    "fundingdesk.com": 4900.00,
    "cranefinancing.com": 5000.00,
    "leasingcapital.com": 6495.00,
    "forgelending.com": 8995.00,
    "gymfinancing.com": 18800.00,
    "anvilcapital.com": 19000.00,
    "leaseflow.com": 70000.00,
    "truckfinancing.com": 375000.00,
    "creditdesk.com": 500000.00,
}
HANDREG = 22.99

wb = load_workbook(WB)
ws = wb["Candidates"]

seen, n_avail = [], 0
for r in range(5, 85):
    dom = ws[f"A{r}"].value
    if not dom:
        continue
    seen.append(dom)
    if dom in AVAILABLE:
        n_avail += 1
        price = AVAILABLE[dom]
        ws[f"K{r}"].value = price
        ws[f"K{r}"].number_format = MONEY2
        status = "Available - reg" if price == HANDREG else "Available - premium"
    else:
        ws[f"K{r}"].value = None
        status = "Taken"
    ws[f"U{r}"].value = status
    for col in ("K", "U"):
        c = ws[f"{col}{r}"]
        c.font = Font(name=FONT, size=10, color=BLUE)
        c.alignment = Alignment(horizontal="center", vertical="center")

missing = [d for d in AVAILABLE if d not in seen]
assert not missing, f"available names not found on the sheet: {missing}"

ws["A2"] = ("Niche 1: equipment financing & commercial lending. Availability and pricing checked at "
            "GoDaddy 2026-08-03: 20 of 56 available (6 at registration price, 14 premium-listed), "
            "36 taken. Premium figures are ASKING prices set by the current holder, not sale "
            "comps — do not feed them back into Est. Resale Value.")
ws["A2"].font = Font(name=FONT, size=10, italic=True)

wb.save(WB)
print(f"checked {len(seen)} rows | available {n_avail} | taken {len(seen)-n_avail}")
